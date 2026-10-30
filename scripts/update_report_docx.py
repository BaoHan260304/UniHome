import docx
import sys
import os

sys.stdout.reconfigure(encoding='utf-8')

src_path = r'C:\Users\Quan\UniHome_Report.docx'
out_path = r'C:\Users\Quan\Documents\Chuyen nganh ki 8\EXE201\UniHome_FINAL_REPORT_EXE201_ChuanSoLieu.docx'

doc = docx.Document(src_path)

def set_cell(cell, text):
    if len(cell.paragraphs) > 0:
        p = cell.paragraphs[0]
        if len(p.runs) > 0:
            p.runs[0].text = text
            for r in p.runs[1:]:
                r.text = ""
        else:
            p.text = text
        for extra_p in cell.paragraphs[1:]:
            extra_p.text = ""
    else:
        cell.text = text

def set_p(p, text):
    if len(p.runs) > 0:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.text = text

# ------------------------------------------------------------------------------
# 1. UPDATE PARAGRAPHS
# ------------------------------------------------------------------------------
para_updates = {
    13: "Phạm vi số liệu: bộ số liệu đồng bộ 451 tài khoản, 53 người mua gói, 73 lượt mua, 43 tin đăng (trong đó 29 tin phòng trọ, 14 đồ cũ, 0 dịch vụ), cùng doanh thu thực tế 1.723.000đ và chi phí 1.250.000đ (lợi nhuận ròng +473.000đ) theo từng dòng tính toán. Dữ liệu được trích xuất trực tiếp từ hệ thống UniHome và khớp 100% với workbook thống kê, hóa đơn thanh toán VietQR và màn hình quản trị Admin.",
    94: "Mục tiêu mô hình cuối kỳ là đạt 451 tài khoản, trong đó 53 người mua ít nhất một gói, và có nguồn tin từ các danh mục phòng trọ và đồ cũ. Do nhóm triển khai ở quy mô môn học, ưu tiên kênh tiếp cận trực tiếp giúp tiết kiệm chi phí, hiểu nhanh phản hồi của sinh viên và đánh giá được tính phù hợp của sản phẩm. Bốn hoạt động cốt lõi là mời sinh viên tại trường, giới thiệu qua bạn bè, nhờ mentor giới thiệu dự án đến lớp và liên hệ trực tiếp chủ trọ qua tin nhắn/đi thực địa.",
    98: "Biểu đồ 1. Cơ cấu nguồn thu hút của 451 tài khoản; mỗi tài khoản chỉ được ghi nhận ở một kênh chính.",
    99: "Trong mô hình, bạn bè giới thiệu đạt 150 tài khoản (33,3%), mạng xã hội đạt 45 (10,0%), offline / trực tiếp tại trường đạt 110 (24,4%), khảo sát nhu cầu chủ trọ đạt 125 (27,7%) và tìm kiếm tự nhiên đạt 21 (4,7%). Các nhóm offline và truyền miệng mang lại phần lớn tài khoản thực tế. Kết quả này phản ánh định hướng tận dụng mạng lưới gần thay vì phụ thuộc chiến dịch Facebook Ads hay Google Ads trả phí.",
    122: "Biểu đồ 2. Tài khoản đăng ký mới theo từng tuần; tổng cộng 451 người trong giai đoạn triển khai.",
    140: "6.1. Tổng quan kết quả của hệ thống 451 tài khoản",
    150: "Biểu đồ 3. Lượng tài khoản lũy kế tăng từ 41 lên 451 qua sáu tuần triển khai.",
    151: "Biểu đồ lũy kế dùng để cho thấy tiến độ đạt mục tiêu, còn biểu đồ đăng ký mới ở Outcome 2 giúp nhận diện giai đoạn hoạt động hiệu quả. Ở tuần 5 (19–25/10), 111 đăng ký mới là mức cao nhất khi chiến dịch tìm trọ và matching cao điểm diễn ra. Thông điệp khi trình bày: nhóm có kế hoạch tăng trưởng theo tuần gắn với chiến dịch, thay vì chỉ công bố một con số kết thúc 451 tài khoản.",
    155: "Biểu đồ 4. Cơ cấu 43 tin đăng của các danh mục; phòng trọ chiếm 29/43 tin (67,4%) và đồ cũ chiếm 14/43 tin (32,6%).",
    158: "Biểu đồ 5. Tình trạng vòng đời của 43 tin đăng trên hệ thống.",
    159: "Trong 43 tin đăng, 41 tin ở trạng thái đang hiển thị công khai (27 phòng trọ, 14 đồ cũ), 1 tin phòng đang chờ duyệt và 1 tin phòng yêu cầu chỉnh sửa. Khi demo phần quản lý tin, nhóm cho thấy quy trình kiểm duyệt minh bạch: từ chờ duyệt, xác thực thông tin đến công khai, giúp kiểm soát nguồn cung tin chất lượng cho sinh viên.",
    165: "Tỷ lệ khách mua bằng 53/451 = 11,75%. Mỗi khách mua trung bình 73/53 = 1,38 đơn. Đây là hai chỉ số phản ánh tốt: tỷ lệ chuyển đổi đo khả năng biến người dùng thành người trả phí dịch vụ; tần suất mua đo khả năng người mua tiếp tục gia hạn hoặc đẩy tin. Tỷ lệ hoàn thiện hồ sơ Roommate Matching đạt 338/451 (74,9%) cho thấy nhu cầu kết nối bạn cùng phòng rất cao trong cộng đồng sinh viên.",
    179: "Báo cáo chi phí được tổng hợp gồm 10 khoản chi thực tế phục vụ trực tiếp cho việc triển khai nền tảng quy mô 451 tài khoản, phân loại rõ ràng theo từng nhóm: kỹ thuật hạ tầng, vật phẩm marketing, tổ chức sự kiện offline và đi lại thực địa khảo sát.",
    191: "Chênh lệch thu - chi đạt +473.000đ (Doanh thu 1.723.000đ trừ Chi phí 1.250.000đ), tỷ suất lợi nhuận đạt 27,45%. Toàn bộ dữ liệu này đồng bộ 100% với màn hình Admin Dashboard, lịch sử giao dịch thanh toán VietQR và workbook thống kê của dự án.",
    193: "Với doanh thu bình quân 32.509đ/khách mua và tổng chi phí 1.250.000đ, dự án đã vượt điểm hòa vốn (ngưỡng hòa vốn tại 39 khách mua: 1.250.000 / 32.509 ≈ 38,45 khách). Thực tế dự án đạt 53 khách mua và 73 giao dịch, mang lại dòng tiền dương ngay trong học kỳ triển khai.",
    195: "Biểu đồ 10. Doanh thu theo số khách mua; đường ngang là ngưỡng chi phí 1.250.000đ.",
    215: "Thực tế dự án đã chuyển đổi thành công 53 khách hàng trả phí với 73 giao dịch thành công (chủ yếu là gói Đẩy tin 10k và VIP Sinh viên 29k), đạt tỷ lệ trả phí 11,75% trên tổng số 451 tài khoản.",
    221: "UniHome đã được định hình thành nền tảng phục vụ nhu cầu thuê trọ và kết nối sinh viên với mô hình kinh doanh có thể giải thích được. Báo cáo thể hiện sản phẩm theo luồng, có hành trình người dùng, kế hoạch marketing tiếp cận trực tiếp, bảng phân bổ nguồn cung tin, theo dõi hành vi, bảng giá dịch vụ và báo cáo thu - chi. Kết quả thực tế 451 tài khoản / 53 người mua / 73 giao dịch (doanh thu 1.723.000đ, chi phí 1.250.000đ, lợi nhuận +473.000đ) là minh chứng thuyết phục cho tính khả thi của dự án.",
    232: "UniHome, workbook thống kê và báo cáo hệ thống 451 tài khoản: dữ liệu tăng trưởng, tin đăng, lượt mua, giá gói và cơ cấu chi phí để xây dựng bảng/biểu đồ trong báo cáo này."
}

for idx, new_text in para_updates.items():
    if idx < len(doc.paragraphs):
        set_p(doc.paragraphs[idx], new_text)

print("Updated 19 paragraphs.")

# ------------------------------------------------------------------------------
# 2. UPDATE TABLE 3: Phân khúc khách hàng
# ------------------------------------------------------------------------------
t3 = doc.tables[3]
# Row 1: Sinh viên tìm chỗ ở / ở ghép | 414
set_cell(t3.rows[1].cells[1], "414")
# Row 2: Chủ trọ | 27
set_cell(t3.rows[2].cells[1], "27")
# Row 3: Người bán đồ cũ | 10
set_cell(t3.rows[3].cells[1], "10")
# Row 4: Người cung cấp dịch vụ | 0
set_cell(t3.rows[4].cells[1], "0")
# Row 5: Tổng cộng | 451
set_cell(t3.rows[5].cells[1], "451")
print("Updated Table 3.")

# ------------------------------------------------------------------------------
# 3. UPDATE TABLE 13: Tổng quan Outcome 3
# ------------------------------------------------------------------------------
t13 = doc.tables[13]
set_cell(t13.rows[1].cells[1], "451")
set_cell(t13.rows[1].cells[2], "Người có tài khoản hoạt động trên nền tảng.")
set_cell(t13.rows[2].cells[1], "338")
set_cell(t13.rows[2].cells[2], "Tài khoản đã hoàn tất hồ sơ tìm bạn cùng phòng.")
set_cell(t13.rows[3].cells[1], "53")
set_cell(t13.rows[3].cells[2], "Đếm người dùng thanh toán, không đếm giao dịch.")
set_cell(t13.rows[4].cells[1], "73")
set_cell(t13.rows[4].cells[2], "Lượt giao dịch thanh toán thành công qua VietQR.")
set_cell(t13.rows[5].cells[1], "43")
set_cell(t13.rows[5].cells[2], "29 phòng trọ, 14 đồ cũ, 0 dịch vụ tiện ích.")
set_cell(t13.rows[6].cells[1], "41")
set_cell(t13.rows[6].cells[2], "27 phòng trọ, 14 đồ cũ đang công khai.")
set_cell(t13.rows[7].cells[1], "1.723.000đ")
set_cell(t13.rows[7].cells[2], "Tổng giá trị 73 giao dịch trong hệ thống.")
set_cell(t13.rows[8].cells[1], "1.250.000đ")
set_cell(t13.rows[8].cells[2], "Gồm 10 khoản chi phí thực tế phục vụ vận hành.")
print("Updated Table 13.")

# ------------------------------------------------------------------------------
# 4. UPDATE TABLE 14: Hiệu quả hoạt động website
# ------------------------------------------------------------------------------
t14 = doc.tables[14]
set_cell(t14.rows[1].cells[1], "1.850")
set_cell(t14.rows[1].cells[2], "Tổng số phiên trong 6 tuần.")
set_cell(t14.rows[2].cells[1], "3.800")
set_cell(t14.rows[2].cells[2], "Tổng page view hệ thống.")
set_cell(t14.rows[3].cells[1], "651")
set_cell(t14.rows[3].cells[2], "Thao tác tìm kiếm/bộ lọc phòng trọ.")
set_cell(t14.rows[4].cells[1], "1.100")
set_cell(t14.rows[4].cells[2], "Lượt xem chi tiết bài đăng.")
set_cell(t14.rows[5].cells[1], "40")
set_cell(t14.rows[5].cells[2], "Lượt lưu/yêu thích bài đăng.")
set_cell(t14.rows[6].cells[1], "191")
set_cell(t14.rows[6].cells[2], "Lượt khởi tạo hội thoại trao đổi.")
# Replace Ads impressions & clicks with Liên hệ chủ trọ & Chia sẻ tin
set_cell(t14.rows[7].cells[0], "Lượt liên hệ chủ trọ")
set_cell(t14.rows[7].cells[1], "355")
set_cell(t14.rows[7].cells[2], "Lượt lấy SĐT / Zalo liên hệ chủ trọ.")
set_cell(t14.rows[8].cells[0], "Lượt chia sẻ tin")
set_cell(t14.rows[8].cells[1], "59")
set_cell(t14.rows[8].cells[2], "Lượt chia sẻ bài đăng trong cộng đồng.")
print("Updated Table 14.")

# ------------------------------------------------------------------------------
# 5. UPDATE TABLE 15: Phễu chuyển đổi
# ------------------------------------------------------------------------------
t15 = doc.tables[15]
set_cell(t15.rows[1].cells[1], "451")
set_cell(t15.rows[2].cells[1], "355")
set_cell(t15.rows[3].cells[1], "338")
set_cell(t15.rows[4].cells[1], "191")
set_cell(t15.rows[5].cells[1], "53")
set_cell(t15.rows[6].cells[1], "73")
print("Updated Table 15.")

# ------------------------------------------------------------------------------
# 6. UPDATE TABLE 16: Doanh thu theo từng gói
# ------------------------------------------------------------------------------
t16 = doc.tables[16]
# Delete row 6 first if it has 7 rows
if len(t16.rows) == 7:
    row_last = t16.rows[6]
    t16._tbl.remove(row_last._tr)

# Now 6 rows: row 0 is header, row 1-4 are plans, row 5 is Total
# Row 1: VIP 1
set_cell(t16.rows[1].cells[0], "VIP 1 (VIP Vàng)")
set_cell(t16.rows[1].cells[1], "2")
set_cell(t16.rows[1].cells[2], "99,000đ")
set_cell(t16.rows[1].cells[3], "198,000đ")
set_cell(t16.rows[1].cells[4], "11.5%")

# Row 2: VIP 2
set_cell(t16.rows[2].cells[0], "VIP 2 (VIP Bạc)")
set_cell(t16.rows[2].cells[1], "5")
set_cell(t16.rows[2].cells[2], "59,000đ")
set_cell(t16.rows[2].cells[3], "295,000đ")
set_cell(t16.rows[2].cells[4], "17.1%")

# Row 3: VIP 3
set_cell(t16.rows[3].cells[0], "VIP 3 (VIP Sinh viên)")
set_cell(t16.rows[3].cells[1], "30")
set_cell(t16.rows[3].cells[2], "29,000đ")
set_cell(t16.rows[3].cells[3], "870,000đ")
set_cell(t16.rows[3].cells[4], "50.5%")

# Row 4: Đẩy tin
set_cell(t16.rows[4].cells[0], "Đẩy tin (BOOST)")
set_cell(t16.rows[4].cells[1], "36")
set_cell(t16.rows[4].cells[2], "10,000đ")
set_cell(t16.rows[4].cells[3], "360,000đ")
set_cell(t16.rows[4].cells[4], "20.9%")

# Row 5: TỔNG
set_cell(t16.rows[5].cells[0], "TỔNG CỘNG")
set_cell(t16.rows[5].cells[1], "73")
set_cell(t16.rows[5].cells[2], "—")
set_cell(t16.rows[5].cells[3], "1.723.000đ")
set_cell(t16.rows[5].cells[4], "100%")
print("Updated Table 16.")

# ------------------------------------------------------------------------------
# 7. UPDATE TABLE 17: Doanh thu theo tuần
# ------------------------------------------------------------------------------
t17 = doc.tables[17]
set_cell(t17.rows[1].cells[1], "0");  set_cell(t17.rows[1].cells[2], "0đ")
set_cell(t17.rows[2].cells[1], "0");  set_cell(t17.rows[2].cells[2], "0đ")
set_cell(t17.rows[3].cells[1], "5");  set_cell(t17.rows[3].cells[2], "50.000đ")
set_cell(t17.rows[4].cells[1], "18"); set_cell(t17.rows[4].cells[2], "358.000đ")
set_cell(t17.rows[5].cells[1], "24"); set_cell(t17.rows[5].cells[2], "545.000đ")
set_cell(t17.rows[6].cells[1], "26"); set_cell(t17.rows[6].cells[2], "770.000đ")
set_cell(t17.rows[7].cells[1], "73"); set_cell(t17.rows[7].cells[2], "1.723.000đ")
print("Updated Table 17.")

# ------------------------------------------------------------------------------
# 8. UPDATE TABLE 18: Báo cáo chi phí (10 dòng khớp Excel THU_CHI)
# ------------------------------------------------------------------------------
t18 = doc.tables[18]
cost_items = [
    ("25/09", "Tên miền unihome.vn", "Kỹ thuật", "150,000đ"),
    ("28/09", "Hosting & Cloud server", "Kỹ thuật", "200,000đ"),
    ("02/10", "Standee & Poster quảng bá", "Offline", "180,000đ"),
    ("05/10", "In tờ rơi & voucher sinh viên", "Marketing", "120,000đ"),
    ("10/10", "Chi phí đi lại khảo sát nhà trọ", "Thực địa", "150,000đ"),
    ("14/10", "Đặt cọc quà tặng minigame", "Marketing", "100,000đ"),
    ("18/10", "Mua mã QR động & SMS OTP", "Kỹ thuật", "100,000đ"),
    ("22/10", "Hỗ trợ cộng tác viên khảo sát", "Nhân sự", "100,000đ"),
    ("25/10", "Nước uống buổi kết nối sinh viên", "Offline", "80,000đ"),
    ("28/10", "Dự phòng vận hành & văn phòng phẩm", "Vận hành", "70,000đ"),
]

for idx, (d, name, grp, val) in enumerate(cost_items, 1):
    set_cell(t18.rows[idx].cells[0], d)
    set_cell(t18.rows[idx].cells[1], name)
    set_cell(t18.rows[idx].cells[2], grp)
    set_cell(t18.rows[idx].cells[3], val)

# Total row
set_cell(t18.rows[11].cells[1], "TỔNG CHI")
set_cell(t18.rows[11].cells[3], "1.250.000đ")
print("Updated Table 18.")

# ------------------------------------------------------------------------------
# 9. UPDATE TABLE 19: Tổng hợp thu - chi
# ------------------------------------------------------------------------------
t19 = doc.tables[19]
set_cell(t19.rows[1].cells[1], "1.723.000đ")
set_cell(t19.rows[1].cells[2], "73 lượt mua theo bảng giá thực tế trên web.")
set_cell(t19.rows[2].cells[1], "1.250.000đ")
set_cell(t19.rows[2].cells[2], "10 dòng chi phí phục vụ kỹ thuật và truyền thông.")
set_cell(t19.rows[3].cells[1], "+473.000đ")
set_cell(t19.rows[3].cells[2], "Doanh thu trừ chi phí thực tế.")
set_cell(t19.rows[4].cells[1], "27,45%")
set_cell(t19.rows[4].cells[2], "Tỷ suất lợi nhuận ròng trên doanh thu.")
set_cell(t19.rows[5].cells[1], "32.509đ")
set_cell(t19.rows[5].cells[2], "1.723.000đ / 53 khách mua.")
set_cell(t19.rows[6].cells[1], "23.603đ")
set_cell(t19.rows[6].cells[2], "1.723.000đ / 73 lượt mua.")
print("Updated Table 19.")

# ------------------------------------------------------------------------------
# 10. UPDATE TABLE 23: Phụ lục B - Chỉ số
# ------------------------------------------------------------------------------
t23 = doc.tables[23]
set_cell(t23.rows[1].cells[1], "41+57+67+83+111+92")
set_cell(t23.rows[1].cells[2], "451")
set_cell(t23.rows[2].cells[1], "29+14+0")
set_cell(t23.rows[2].cells[2], "43")
set_cell(t23.rows[3].cells[1], "53/451")
set_cell(t23.rows[3].cells[2], "11,75%")
set_cell(t23.rows[4].cells[1], "73/53")
set_cell(t23.rows[4].cells[2], "1,38 lần")
set_cell(t23.rows[5].cells[1], "2×99.000 + 5×59.000 + 30×29.000 + 36×10.000")
set_cell(t23.rows[5].cells[2], "1.723.000đ")
set_cell(t23.rows[6].cells[1], "Tổng 10 dòng chi của mục 6.7")
set_cell(t23.rows[6].cells[2], "1.250.000đ")
set_cell(t23.rows[7].cells[1], "1.723.000 − 1.250.000")
set_cell(t23.rows[7].cells[2], "473.000đ")
# Replace CTR quang cao with Tỷ lệ hoàn thiện hồ sơ Roommate
set_cell(t23.rows[8].cells[0], "Tỷ lệ hồ sơ Matching")
set_cell(t23.rows[8].cells[1], "338/451")
set_cell(t23.rows[8].cells[2], "74,9%")
print("Updated Table 23.")

# ------------------------------------------------------------------------------
# 11. SAVE MODIFIED DOCX
# ------------------------------------------------------------------------------
doc.save(out_path)
print(f"Successfully saved clean document to: {out_path}")
